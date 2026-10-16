"""Export a question + grounded answer + cited excerpts as a legal-memo .docx."""
from __future__ import annotations

import io
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Sequence

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from config import DISCLAIMER  # noqa: E402
from retrieval.search import Hit  # noqa: E402

_PLAIN_DISCLAIMER = re.sub(r"\*\*", "", DISCLAIMER)


def _rule(doc):
    from docx.shared import Pt
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    pr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "888888")
    bdr.append(bottom)
    pr.append(bdr)


def build_memo(question: str, answer: str, hits: Sequence[Hit],
               model: str = "", filters_desc: str = "") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)

    # ---- heading
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("RESEARCH MEMORANDUM")
    r.bold = True
    r.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("LegalCase Responder — property-law corpus")
    sr.italic = True
    sr.font.size = Pt(10)
    sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    _rule(doc)

    # ---- header block
    meta = doc.add_paragraph()
    meta.add_run("Date:\t").bold = True
    meta.add_run(datetime.now().strftime("%B %d, %Y  %H:%M"))
    meta.add_run("\nPrepared by:\t").bold = True
    meta.add_run(f"LegalCase Responder (RAG){f' — {model}' if model else ''}")
    meta.add_run("\nRetrieval:\t").bold = True
    meta.add_run(f"{len(hits)} excerpt(s){f'; filters: {filters_desc}' if filters_desc else ''}")

    # ---- question
    doc.add_heading("Question Presented", level=1)
    doc.add_paragraph(question.strip())

    # ---- answer
    doc.add_heading("Answer", level=1)
    for para in [p for p in (answer or "").split("\n") if p.strip()]:
        doc.add_paragraph(para.strip())
    if not (answer or "").strip():
        doc.add_paragraph("(No answer was generated.)")

    # ---- sources
    doc.add_heading("Sources", level=1)
    intro = doc.add_paragraph()
    ir = intro.add_run("Each excerpt below is reproduced verbatim from the corpus and is "
                       "the only material the answer was permitted to rely on.")
    ir.italic = True
    ir.font.size = Pt(9)

    for i, hit in enumerate(hits, start=1):
        m = hit.meta
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        cr = p.add_run(f"{i}. {hit.citation()}")
        cr.bold = True
        tag = doc.add_paragraph()
        tr = tag.add_run(
            f"{m.get('court','')} · {m.get('jurisdiction','')} · {m.get('date','')} · "
            f"case_type: {m.get('case_type','')} · source: {hit.source.upper()} · "
            f"case_id: {m.get('case_id','')} · similarity {hit.score:.3f}")
        tr.font.size = Pt(8)
        tr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if m.get("heading"):
            hp = doc.add_paragraph()
            hr = hp.add_run(m["heading"])
            hr.italic = True
            hr.font.size = Pt(9)
        body = doc.add_paragraph(hit.text)
        body.paragraph_format.left_indent = Inches(0.35)
        body.paragraph_format.space_after = Pt(4)
        for run in body.runs:
            run.font.size = Pt(10)

    if not hits:
        doc.add_paragraph("(No excerpts were retrieved.)")

    # ---- disclaimer
    _rule(doc)
    d = doc.add_paragraph()
    dr = d.add_run(_PLAIN_DISCLAIMER)
    dr.italic = True
    dr.font.size = Pt(9)
    dr.font.color.rgb = RGBColor(0x99, 0x33, 0x33)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def suggested_filename(question: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", question.lower()).strip("-")[:48] or "memo"
    return f"legalcase-memo-{slug}-{datetime.now():%Y%m%d-%H%M}.docx"
